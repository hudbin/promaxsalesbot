# -*- coding: utf-8 -*-
"""Kunlik hisobot: Telegram matni (HTML) va Excel fayl."""
import re
from datetime import date
from html import escape as _escape


def escape(s):
    """HTML uchun xavfsiz matn (apostrof o'zgarmaydi: og'a)."""
    return _escape(str(s), quote=False)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import ainur
import config
import daftar
import db

HAFTA = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]


# --- Yordamchi funksiyalar --------------------------------------------------------
def pul(x):
    """1500000 → '1 500 000'."""
    return f"{int(round(x or 0)):,}".replace(",", " ").replace("-", "−")


def summa_oqi(matn):
    """
    Foydalanuvchi yozgan summani o'qish:
    '1500000', '1 500 000', '1.500.000', '1,5 mln', '1.5m', '500 ming', '500k'.
    Noto'g'ri bo'lsa None.
    """
    s = str(matn).lower().replace("\u00a0", "").replace(" ", "")
    for v in ("so'm", "som", "sum", "сум", "uzs"):
        s = s.replace(v, "")
    kop = 1
    for suf, k in (("million", 10**6), ("mln", 10**6), ("млн", 10**6), ("m", 10**6),
                   ("ming", 10**3), ("тыс", 10**3), ("k", 10**3)):
        if s.endswith(suf):
            s, kop = s[: -len(suf)], k
            break
    if kop == 1:
        s = s.replace(",", "").replace(".", "")   # ajratkichlar
    else:
        s = s.replace(",", ".")                   # 1,5 mln → 1.5
    try:
        v = round(float(s) * kop)
    except ValueError:
        return None
    return v if 0 < v < 10**11 else None


def sana_oqi(matn, bugun):
    """Izohdagi sanani o'qish: '24.09', '24.09.2026', '24/09/26'. Topilmasa None."""
    m = re.search(r"\b(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?\b", matn or "")
    if not m:
        return None
    k, o, y = int(m.group(1)), int(m.group(2)), m.group(3)
    y = int(y) + (2000 if y and len(y) == 2 else 0) if y else bugun.year
    try:
        return date(y, o, k)
    except ValueError:
        return None


def sana_matn(d):
    return f"{d:%d.%m.%Y} ({HAFTA[d.weekday()]})"


def _b(s):
    return f"<b>{escape(str(s))}</b>"


# --- Narx nazorati --------------------------------------------------------------
def narx_muammolari(qatorlar):
    """0 so'mga yoki tannarxdan past sotilgan pozitsiyalar."""
    nol = [q for q in qatorlar if q["sotilgan"] > 0 and q["daromad"] == 0]
    past = [q for q in qatorlar
            if q["daromad"] > 0 and q["tannarx"] > 0 and q["daromad"] < q["tannarx"] * 0.95]
    nol.sort(key=lambda q: -q["sotilgan"])
    past.sort(key=lambda q: q["daromad"] - q["tannarx"])
    return nol, past


def fayl_xulosa(tur, qatorlar, sana):
    """Excel qabul qilingandan keyin yuklovchiga qisqa javob."""
    if tur == "sotuv":
        dona = sum(q["sotilgan"] for q in qatorlar)
        summa = sum(q["daromad"] for q in qatorlar)
        nol, past = narx_muammolari(qatorlar)
        s = (f"✅ Savdo hisoboti qabul qilindi — {sana:%d.%m.%Y}\n"
             f"{len(qatorlar)} pozitsiya · {pul(dona)} dona · {pul(summa)} {config.VALYUTA}")
        if nol:
            s += f"\n⚠️ 0 so'mga chiqqan: {len(nol)} pozitsiya, {pul(sum(q['sotilgan'] for q in nol))} dona"
        if past:
            s += f"\n⚠️ Tannarxdan past: {len(past)} pozitsiya"
        return s
    manfiy = [q for q in qatorlar if q["yakuniy"] < 0]
    s = (f"✅ Ombor hisoboti qabul qilindi — {sana:%d.%m.%Y}\n"
         f"{len(qatorlar)} pozitsiya · qoldiq {pul(sum(q['yakuniy'] for q in qatorlar))} dona")
    if manfiy:
        s += f"\n⚠️ Manfiy qoldiq: {len(manfiy)} pozitsiya"
    return s


# --- Telegram matni ---------------------------------------------------------------
# Tartib: QISQACHA → MODELLAR → MIJOZLAR → ESKI QARZ TO'LOVLARI → QAYTARISH →
#         NARXSIZ/ICHKI CHIQIM → QARZDORLAR → OMBOR
def _n(s):
    """Nomni solishtirish uchun normallashtirish."""
    s = (s or "").lower()
    for a in "ʻʼ‘’`":
        s = s.replace(a, "'")
    return re.sub(r"\s+", " ", s).strip()


_ICHKI = {_n(x) for x in getattr(config, "ICHKI_MIJOZLAR", ["X"])}


def _turlar_matn(turlar):
    return ", ".join(f"{t} {pul(v)}" for t, v in turlar.items() if v)


def _qisqacha(q, hj, qt, qarz, qol=(), rx=()):
    savdo = sum(x["daromad"] for x in q)
    dona = sum(x["sotilgan"] for x in q)
    soni = len(hj) if hj else int(sum(x["sotish"] for x in q))
    kassa = sum(h["tolangan"] for h in hj)
    qarzga = sum(max(h["jami"] - h["tolangan"], 0) for h in hj)
    eski = sum(t["summa"] for t in qt) + sum(t["summa"] for t in qol)
    turlar = {}
    for x in list(hj) + list(qt) + [{"turlar": {t["turi"]: t["summa"]}} for t in qol]:
        for t, v in x["turlar"].items():
            turlar[t] = turlar.get(t, 0) + v
    L = ["📌 <b>QISQACHA</b>",
         f"🛒 Sotuv: <b>{pul(savdo)} {config.VALYUTA}</b> · {pul(dona)} dona · {soni} hujjat"]
    if hj:
        L += [f"💵 Savdodan to'landi: {pul(kassa)}", f"📝 Qarzga berildi: <b>{pul(qarzga)}</b>"]
    L.append(f"💳 Eski qarzdan tushdi: {pul(eski)}")
    L.append(f"💰 Jami tushum: <b>{pul(kassa + eski)}</b>"
             + (f" ({_turlar_matn(turlar)})" if turlar else ""))
    if rx:
        rasxod = sum(x["summa"] for x in rx)
        kat = {}
        for x in rx:
            kat[x["kategoriya"]] = kat.get(x["kategoriya"], 0) + x["summa"]
        ichi = ", ".join(f"{escape(k)} {pul(v)}" for k, v in
                         sorted(kat.items(), key=lambda i: -i[1]))
        L.append(f"🧾 Rasxod: <b>{pul(rasxod)}</b> ({ichi})")
        L.append(f"📈 Sof foyda: <b>{pul(savdo - rasxod)} {config.VALYUTA}</b>")
    if qarz:
        farq = "" if qarz["farq"] is None else \
            f" ({'+' if qarz['farq'] >= 0 else '−'}{pul(abs(qarz['farq']))})"
        L.append(f"📒 Umumiy qarzdorlik: <b>{pul(qarz['jami'])}</b>{farq}")
    return L + [""]


def _modellar(q):
    guruh = [g for g in ainur.modellarga_guruhla(q) if g["dona"] or g["summa"]]
    if not guruh:
        return []
    L = ["🛒 <b>MODELLAR</b>"]
    for i, g in enumerate(guruh, 1):
        narx = g["summa"] / g["dona"] if g["dona"] else 0
        L.append(f"{i}. {_b(g['model'])} — {pul(g['dona'])} dona × {pul(narx)} = {pul(g['summa'])}")
        if len(g["narxlar"]) > 1:  # bir model turli narxlarda sotilgan
            tar = sorted(g["narxlar"].items(), key=lambda x: -x[1])[:4]
            L.append("    <i>narxlar: " + ", ".join(f"{pul(n)}×{pul(d)}" for n, d in tar) + "</i>")
    return L + [""]


def _mijozlar(hj):
    if not hj:
        return []
    mij = {}
    for h in hj:
        m = mij.setdefault(h["mijoz"] or "Mijozsiz",
                           {"jami": 0, "tol": 0, "turlar": {}, "xodim": [], "soni": 0})
        m["jami"] += h["jami"]
        m["tol"] += h["tolangan"]
        m["soni"] += 1
        for t, v in h["turlar"].items():
            m["turlar"][t] = m["turlar"].get(t, 0) + v
        if h["xodim"] and h["xodim"] not in m["xodim"]:
            m["xodim"].append(h["xodim"])
    L = ["👥 <b>MIJOZLAR</b>"]
    for i, (nom, m) in enumerate(sorted(mij.items(), key=lambda x: -x[1]["jami"]), 1):
        qism = [f"{i}. {_b(nom)} — {pul(m['jami'])}"]
        if m["tol"]:
            tt = _turlar_matn(m["turlar"])
            qism.append(f"to'landi {pul(m['tol'])}" + (f" ({tt})" if tt else ""))
        qarz = max(m["jami"] - m["tol"], 0)
        if qarz:
            qism.append(f"qarzga <b>{pul(qarz)}</b>")
        if m["xodim"]:
            qism.append(escape(", ".join(m["xodim"])))
        L.append(" · ".join(qism))
    return L + [""]


def _rasxodlar(rx):
    """Rasxodlar kategoriya bo'yicha, har birining izohi bilan."""
    if not rx:
        return []
    kat = {}
    for x in rx:
        kat.setdefault(x["kategoriya"], []).append(x)
    L = [f"🧾 <b>RASXODLAR</b> ({len(rx)} ta · {pul(sum(x['summa'] for x in rx))} {config.VALYUTA})"]
    for nom, qator in sorted(kat.items(), key=lambda i: -sum(x["summa"] for x in i[1])):
        L.append(f"📂 <b>{escape(nom)}</b> — {pul(sum(x['summa'] for x in qator))}")
        for x in qator:
            sana = f"{x['sana'][8:10]}.{x['sana'][5:7]} " if len({y['sana'] for y in rx}) > 1 else ""
            L.append(f"  • {pul(x['summa'])} · {sana}{x['vaqt']} · {escape(x['kim'] or '')}")
            if x["izoh"]:
                L.append(f"    💬 {escape(x['izoh'])}")
    return L + [""]


def _qol_tolovlar(qol):
    """Mini app'da qayd qilingan qarz to'lovlari (AinurPOS'ga qo'lda kiritiladi)."""
    if not qol:
        return []
    L = [f"📝 <b>QARZ TO'LOVLARI — mini app</b> ({len(qol)} ta)"]
    for i, t in enumerate(qol, 1):
        belgi = "✅" if t["ainurga"] else "⚠️ AinurPOS'ga kiritilmagan"
        L.append(f"{i}. {_b(t['mijoz'] or 'Mijozsiz')} — <b>{pul(t['summa'])}</b> "
                 f"({escape(t['turi'])}) · {escape(t['kim'])} · {belgi}")
    L.append(f"<b>Jami: {pul(sum(t['summa'] for t in qol))} {config.VALYUTA}</b>")
    return L + [""]


def _qarz_tolovlari(qt):
    if not qt:
        return []
    L = [f"💳 <b>ESKI QARZ TO'LOVLARI</b> ({len(qt)} ta)"]
    for i, t in enumerate(sorted(qt, key=lambda x: -x["summa"]), 1):
        tt = _turlar_matn(t["turlar"])
        L.append(f"{i}. {_b(t['mijoz'] or 'Mijozsiz')} — <b>{pul(t['summa'])}</b>"
                 + (f" ({tt})" if tt else "") + f" · hujjat #{escape(t['raqam'])}, {_dm(t['hujjat_sana'])}")
    return L + [""]


def _qaytarish(qayt):
    if not qayt:
        return []
    L = [f"↩️ <b>QAYTARISH</b>: {len(qayt)} ta · {pul(sum(x['dona'] for x in qayt))} dona · "
         f"{pul(sum(x['jami'] for x in qayt))}"]
    L += [f"{i}. {escape(x['mijoz'] or 'Mijozsiz')} — {pul(x['dona'])} dona · {pul(x['jami'])}"
          for i, x in enumerate(qayt[:10], 1)]
    return L + [""]


def _narxsiz(qatorlar, hujjatlar):
    """0 so'mlik yoki ichki («X») chiqimlar — savdoga qo'shilmaydi, alohida ko'rsatiladi."""
    qatorlar = [x for x in qatorlar if x["sotilgan"]]
    if not qatorlar:
        return []
    L = ["⚠️ <b>NARXSIZ / ICHKI CHIQIM</b> <i>(savdoga qo'shilmadi)</i>"]
    kimlar = {}
    for h in hujjatlar:
        kimlar[h["mijoz"] or "Mijozsiz"] = kimlar.get(h["mijoz"] or "Mijozsiz", 0) + 1
    kim = ", ".join(f"«{escape(k)}» {v} hujjat" for k, v in kimlar.items()) or "0 so'mga sotilgan"
    L.append(f"{kim} · jami <b>{pul(sum(x['sotilgan'] for x in qatorlar))} dona</b>")
    gr = ainur.modellarga_guruhla(qatorlar)
    L.append("Modellar: " + ", ".join(f"{escape(g['model'])} {pul(g['dona'])}" for g in
                                      sorted(gr, key=lambda g: -g["dona"])[:12])
             + (" …" if len(gr) > 12 else ""))
    L.append("<i>Savdo bo'lsa — AinurPOS'da narx va mijozni to'g'rilang.</i>")
    return L + [""]


def _tannarx_ogoh(q):
    _, past = narx_muammolari(q)
    if not past:
        return []
    top = ", ".join(f"{escape(x['nomi'])} (o'rt. {pul(x['daromad'] / x['sotilgan'])}, "
                    f"tannarx {pul(x['tannarx'] / x['sotilgan'])})" for x in past[:3] if x["sotilgan"])
    return [f"⚠️ <b>Tannarxdan past sotilgan</b>: {len(past)} poz. — {top}", ""]


# Hujjat turlari: belgi, nomi, kontragent yorlig'i
TUR = {
    "kirim": ("🟢", "KIRIM", "🏭 Kimdan"),
    "qayt_mijoz": ("↩️", "QAYTARISH · mijozdan", "👤 Mijoz"),
    "sotuv": ("🔴", "CHIQIM · SOTUV", "👤 Mijoz"),
    "ichki": ("⚪", "ICHKI CHIQIM (0 so'm)", "👤 Kimga"),
    "qayt_yetk": ("↪️", "QAYTARISH · yetkazib beruvchiga", "🏭 Kimga"),
    "tuzatish": ("🔧", "QOLDIQ TUZATISH", ""),
}


def harakat_matn(h, ochirildi=False):
    """Lenta xabari: bitta AinurPOS hujjati (yangi yoki o'chirilgan)."""
    kalit = "ichki" if h.get("ichki") else h["tur"]
    belgi, nom, yorliq = TUR[kalit]
    sarl = f"#{escape(h['raqam'])} · {_dm(h['sana'])} {h['vaqt']}"
    L = [f"🗑 <b>O'CHIRILDI</b> — {nom} {sarl}" if ochirildi else f"{belgi} <b>{nom}</b> {sarl}"]
    kim = []
    if h.get("kontragent") and yorliq:
        kim.append(f"{yorliq}: {escape(h['kontragent'])}")
    if h.get("xodim"):
        kim.append(f"✍️ {escape(h['xodim'])}")
    if kim:
        L.append(" · ".join(kim))
    if h.get("izoh"):
        L.append(f"💬 {escape(h['izoh'][:120])}")
    qat = h.get("qatorlar") or []
    for x in qat[:15]:
        if kalit == "tuzatish":
            L.append(f"• {escape(x['nomi'])} — {'+' if x['soni'] > 0 else '−'}{pul(abs(x['soni']))}")
        elif x["narx"]:
            L.append(f"• {escape(x['nomi'])} — {pul(x['soni'])} × {pul(x['narx'])} = "
                     f"{pul(x['soni'] * x['narx'])}")
        else:
            L.append(f"• {escape(x['nomi'])} — {pul(x['soni'])} dona")
    if len(qat) > 15:
        L.append(f"<i>… yana {len(qat) - 15} pozitsiya</i>")
    if kalit == "tuzatish":
        plus = sum(x["soni"] for x in qat if x["soni"] > 0)
        minus = -sum(x["soni"] for x in qat if x["soni"] < 0)
        L.append(f"<b>Jami: +{pul(plus)} / −{pul(minus)} dona</b>")
    else:
        L.append(f"<b>Jami: {pul(h['dona'])} dona" +
                 (f" · {pul(h['summa'])} {config.VALYUTA}" if h["summa"] else "") + "</b>")
    if kalit == "sotuv" and not ochirildi:
        L.append(f"💵 To'landi: {pul(h['tolangan'])} · 📝 Qarzga: "
                 f"<b>{pul(max(h['summa'] - h['tolangan'], 0))}</b>")
    return "\n".join(L)


def _kirim_chiqim(harakat):
    """Kun/davr bo'yicha barcha AinurPOS hujjatlari turlar kesimida (lenta ma'lumoti)."""
    if not harakat:
        return []
    L = ["📦 <b>KIRIM-CHIQIM</b> <i>(barcha AinurPOS hujjatlari)</i>"]
    faol = [h for h in harakat if not h["ochirilgan"]]
    for kalit, (belgi, nom, _) in TUR.items():
        hh = [h for h in faol if ("ichki" if h["ichki"] else h["tur"]) == kalit]
        if not hh:
            continue
        dona = sum(h["dona"] for h in hh)
        dona_m = (f"{'+' if dona >= 0 else '−'}{pul(abs(dona))}" if kalit == "tuzatish" else pul(dona))
        summa = sum(h["summa"] for h in hh)
        L.append(f"{belgi} {nom.capitalize()}: {len(hh)} hujjat · {dona_m} dona"
                 + (f" · {pul(summa)}" if summa else ""))
    och = [h for h in harakat if h["ochirilgan"]]
    if och:
        L.append(f"🗑 O'chirilgan: {len(och)} hujjat (" + ", ".join(
            f"#{escape(h['raqam'])}" for h in och[:8]) + ")")
    return L + [""]


def _qarz_holati(sana):
    """Qarzdorlik surati: jami, soni, oldingi kunga nisbatan farq, eng katta qarzdorlar."""
    bugun = [m for m in db.qarz_ol(sana.isoformat()) if _n(m["nom"]) not in _ICHKI]
    if not bugun:
        return None
    qarzdor = sorted([m for m in bugun if m["qarz"] >= 1], key=lambda m: -m["qarz"])
    jami = sum(m["qarz"] for m in qarzdor)
    old = db.oldingi_qarz_sanasi(sana.isoformat())
    farq = None
    if old:
        farq = jami - sum(max(m["qarz"], 0) for m in db.qarz_ol(old) if _n(m["nom"]) not in _ICHKI)
    return {"jami": jami, "soni": len(qarzdor), "farq": farq, "top": qarzdor[:10]}


def _qarzdorlar(qarz):
    if not qarz or not qarz["top"]:
        return []
    L = [f"📒 <b>ENG KATTA QARZDORLAR</b> (jami {pul(qarz['jami'])}, {qarz['soni']} mijoz)"]
    L += [f"{i}. {escape(m['nom'])} — {pul(m['qarz'])}" for i, m in enumerate(qarz["top"], 1)]
    return L + [""]


def _asosiy(q, hj, qayt, iq, ih, qt, qarz=None, harakat=None, qol=(), rx=()):
    """Hisobotning asosiy qismi (kunlik va davr uchun umumiy)."""
    # Excel rejimida 0 so'mlik qatorlar ham narxsiz blokka o'tadi
    nol = [x for x in q if x["sotilgan"] > 0 and x["daromad"] == 0]
    q = [x for x in q if not (x["sotilgan"] > 0 and x["daromad"] == 0)]
    L = _qisqacha(q, hj, qt, qarz, qol, rx)
    L += _modellar(q)
    L += _mijozlar(hj)
    L += _qarz_tolovlari(qt)
    L += _qol_tolovlar(qol)
    L += _rasxodlar(rx)
    L += _qaytarish(qayt)
    L += _narxsiz(list(iq) + nol, ih)
    L += _tannarx_ogoh(q)
    L += _kirim_chiqim(harakat or [])
    L += _qarzdorlar(qarz)
    return L


def _daftar_malumot(bosh, oxir):
    """Daftar rejimi: savdo, hujjatlar, to'lovlar va qarzdorlik botning o'z bazasidan."""
    b, o = bosh.isoformat(), oxir.isoformat()
    q = daftar.savdo_qatorlari(b, o)
    hj = [{"raqam": str(x["id"]), "vaqt": x["vaqt"] if bosh == oxir
           else f"{x['sana'][8:10]}.{x['sana'][5:7]} {x['vaqt']}",
           "mijoz": x["mijoz"] or "Chakana", "xodim": x["kim"] or "",
           "jami": x["jami"], "tolangan": x["tolangan"],
           "turlar": {x["turi"]: x["tolangan"]} if x["turi"] and x["tolangan"] else {}}
          for x in daftar.savdolar(b, o)]
    qt = [{"mijoz": t["mijoz"], "summa": t["summa"], "turlar": {t["turi"]: t["summa"]},
           "raqam": str(t["id"]), "hujjat_sana": t["sana"], "sana": t["sana"]}
          for t in daftar.tolovlar(b, o)]
    mij = daftar.mijozlar(faqat_qarzdor=True)
    qarz = {"jami": sum(m["qarz"] for m in mij), "soni": len(mij), "farq": None,
            "top": [{"nom": m["nom"], "qarz": m["qarz"]} for m in mij[:10]]}
    return q, hj, qt, qarz


def kunlik_matn(sana):
    """Kunlik hisobot matni (HTML qatorlar ro'yxati)."""
    s = sana.isoformat()
    if daftar.kochirildimi():
        q, hj, qt, qarz = _daftar_malumot(sana, sana)
        L = ["📊 <b>PROMAX · KUNLIK HISOBOT</b>",
             f"📅 {sana_matn(sana)} · <i>bot daftari</i>", ""]
        L += _asosiy(q, hj, [], [], [], qt, qarz, db.kun_harakatlari(s), (), db.rasxodlar(s))
        navbat = daftar.navbat_soni()
        if navbat:
            L.append(f"⚠️ AinurPOS omboriga yuborilmagan savdo: <b>{navbat}</b> ta")
        return L
    f = db.oxirgi_fayl(s, "sotuv")
    fi = db.oxirgi_fayl(s, "ichki")
    manba = "ma'lumot yo'q"
    if f:
        manba = ("AinurPOS" if f["yukladi"] == "AinurPOS API" else f"Excel · {f['yukladi']}") + \
            f", {f['yuklandi'][11:16]}"
    L = ["📊 <b>PROMAX · KUNLIK HISOBOT</b>", f"📅 {sana_matn(sana)} · <i>{escape(manba)}</i>", ""]
    if not f:
        L += ["⚠️ Bu kun uchun savdo ma'lumoti yo'q.", ""]
    L += _asosiy(db.fayl_qatorlari(f) if f else [], db.fayl_hujjatlari(f) if f else [],
                 db.fayl_qaytarishlari(f) if f else [], db.fayl_qatorlari(fi) if fi else [],
                 db.fayl_hujjatlari(fi) if fi else [], db.kun_qarz_tolovlari(s), _qarz_holati(sana),
                 db.kun_harakatlari(s), db.qol_tolovlar(s), db.rasxodlar(s))

    # OMBOR — ixtiyoriy (qoldiq Excel yuklangan bo'lsa)
    f = db.oxirgi_fayl(s, "qoldiq")
    if f:
        q = db.fayl_qatorlari(f)
        L.append(f"📦 <b>OMBOR</b> <i>(Excel, {f['yuklandi'][11:16]})</i>")
        L.append(f"Kelish: {pul(sum(x['kelish'] for x in q))} · "
                 f"Chiqim: {pul(sum(x['istemol'] for x in q))} · "
                 f"Qoldiq: <b>{pul(sum(x['yakuniy'] for x in q))}</b> dona")
        manfiy = [x for x in q if x["yakuniy"] < 0]
        if manfiy:
            L.append("⚠️ Manfiy qoldiq: " + ", ".join(
                f"{escape(x['nomi'])} ({pul(x['yakuniy'])})" for x in manfiy[:10]))
    return L


# --- Davr hisoboti (kecha / hafta / oy) ------------------------------------------------
def davr_matn(bosh, oxir, q, hj, qayt, iq=(), ih=()):
    L = ["📈 <b>PROMAX · DAVR HISOBOTI</b>",
         f"📅 {bosh:%d.%m.%Y} — {oxir:%d.%m.%Y} ({(oxir - bosh).days + 1} kun)", ""]
    if daftar.kochirildimi():
        dq, dhj, dqt, dqarz = _daftar_malumot(bosh, oxir)
        return L + _asosiy(dq, dhj, [], [], [], dqt, dqarz,
                           db.davr_harakatlari(bosh.isoformat(), oxir.isoformat()), (),
                           db.rasxodlar(bosh.isoformat(), oxir.isoformat()))
    return L + _asosiy(q, hj, qayt, iq, ih,
                       db.davr_qarz_tolovlari(bosh.isoformat(), oxir.isoformat()),
                       harakat=db.davr_harakatlari(bosh.isoformat(), oxir.isoformat()),
                       qol=db.qol_tolovlar(bosh.isoformat(), oxir.isoformat()),
                       rx=db.rasxodlar(bosh.isoformat(), oxir.isoformat()))


def davr_excel(bosh, oxir, q, hj, yol):
    wb = Workbook()
    wb.remove(wb.active)
    gr = ainur.modellarga_guruhla(q)
    _varaq(wb, "Modellar", ["Model", "Dona", "O'rtacha narx", "Summa"],
           [[g["model"], g["dona"], round(g["summa"] / g["dona"]) if g["dona"] else 0, g["summa"]]
            for g in gr], [28, 10, 14, 16],
           ["JAMI", sum(g["dona"] for g in gr), None, sum(g["summa"] for g in gr)])
    _varaq(wb, "Hujjatlar", ["Vaqt", "Raqam", "Mijoz", "Sotuvchi", "Jami", "To'landi", "Qarzga"],
           [[h["vaqt"], h["raqam"], h["mijoz"], h["xodim"], h["jami"], h["tolangan"],
             max(h["jami"] - h["tolangan"], 0)] for h in hj],
           [12, 10, 26, 16, 16, 16, 16],
           ["JAMI", None, None, None, sum(h["jami"] for h in hj), sum(h["tolangan"] for h in hj),
            sum(max(h["jami"] - h["tolangan"], 0) for h in hj)])
    wb.save(yol)
    return yol


# --- Qarzdorlik va mijozlar (menyu uchun) ----------------------------------------------
def qarzdorlik_matn(royxat, yangilangan, limit=30):
    """royxat: [{nom, qarz, ochiq, oxirgi}] — daftar yoki AinurPOS'dan."""
    qarzdor = [m for m in royxat if (m["qarz"] or 0) >= 1]
    jami = sum(m["qarz"] for m in qarzdor)
    L = ["📒 <b>UMUMIY QARZDORLIK</b>",
         f"<i>AinurPOS savdo hujjatlari bo'yicha · yangilangan {yangilangan}</i>", "",
         f"Jami qarz: <b>{pul(jami)} {config.VALYUTA}</b>",
         f"Qarzdor mijozlar: <b>{len(qarzdor)}</b> / {len(royxat)}", ""]
    for i, m in enumerate(qarzdor[:limit], 1):
        L.append(f"{i}. {escape(m['nom'])} — <b>{pul(m['qarz'])}</b> "
                 f"<i>({m['ochiq']} hujjat, oxirgi {_dm(m['oxirgi'])})</i>")
    if len(qarzdor) > limit:
        L.append(f"… yana {len(qarzdor) - limit} mijoz — to'liq ro'yxat Excel'da")
    return L


def _dm(iso):
    return f"{iso[8:10]}.{iso[5:7]}" if iso else "—"


def mijoz_karta_matn(m, qarz, hujjatlar, tolovlar):
    """Mijoz kartasi: qarz, xaridlar, ochiq hujjatlar, oxirgi qarz to'lovlari."""
    L = [f"👤 <b>{escape(m['nom'])}</b>"]
    if m.get("telefon"):
        L.append(f"📞 {escape(m['telefon'])}")
    L += ["", f"Qarz: <b>{pul(qarz)} {config.VALYUTA}</b>"]
    if hujjatlar:
        L.append(f"Xaridlar: {len(hujjatlar)} ta hujjat · {pul(sum(h['dona'] or 0 for h in hujjatlar))} dona · "
                 f"{pul(sum(h['jami'] for h in hujjatlar))} {config.VALYUTA} "
                 f"({_dm(hujjatlar[-1]['sana'])} — {_dm(hujjatlar[0]['sana'])})")
        ochiq = [h for h in hujjatlar if h["jami"] - h["tolangan"] >= 1]
        if ochiq:
            L += ["", "<b>To'lanmagan hujjatlar:</b>"]
            L += [f"• #{escape(h['raqam'])} {_dm(h['sana'])} — {pul(h['jami'])}, to'landi "
                  f"{pul(h['tolangan'])}, <b>qoldiq {pul(h['jami'] - h['tolangan'])}</b>"
                  for h in ochiq[:10]]
        L += ["", "<b>Oxirgi xaridlar:</b>"]
        L += [f"• #{escape(h['raqam'])} {_dm(h['sana'])} — {pul(h['dona'] or 0)} dona · {pul(h['jami'])}"
              for h in hujjatlar[:5]]
    else:
        L.append("AinurPOS'da xarid hujjatlari yo'q.")
    if tolovlar:
        L += ["", "<b>Oxirgi qarz to'lovlari:</b>"]
        for t in tolovlar:
            turi = ", ".join(f"{k} {pul(v)}" for k, v in t["turlar"].items()) or "turi ko'rsatilmagan"
            L.append(f"• {_dm(t['sana'])} — {pul(t['summa'])} ({turi})")
    return L


def qarz_excel(royxat, yol):
    wb = Workbook()
    wb.remove(wb.active)
    qarzdor = [m for m in royxat if (m["qarz"] or 0) >= 1]
    _varaq(wb, "Qarzdorlar", ["Mijoz", "Telefon", "Qarz", "Ochiq hujjat", "Oxirgi xarid"],
           [[m["nom"], m["telefon"], m["qarz"], m["ochiq"], m["oxirgi"]] for m in qarzdor],
           [30, 16, 16, 12, 12], ["JAMI", None, sum(m["qarz"] for m in qarzdor), None, None])
    ochiq = []
    for m in qarzdor:
        for h in db.mijoz_hujjatlari(m["id"]):
            if h["jami"] - h["tolangan"] >= 1:
                ochiq.append([m["nom"], h["raqam"], h["sana"], h["jami"], h["tolangan"],
                              h["jami"] - h["tolangan"]])
    _varaq(wb, "OchiqHujjatlar", ["Mijoz", "Hujjat", "Sana", "Jami", "To'landi", "Qoldiq"],
           ochiq, [30, 10, 12, 16, 16, 16],
           ["JAMI", None, None, None, None, sum(x[5] for x in ochiq)])
    _varaq(wb, "Mijozlar", ["Mijoz", "Telefon", "Hujjatlar", "Qarz"],
           [[m["nom"], m["telefon"], m["hujjat"], m["qarz"]] for m in royxat], [30, 16, 10, 16])
    wb.save(yol)
    return yol


def bolaklarga(qatorlar, limit=3900):
    """Telegram 4096 belgi chegarasi uchun qatorlarni bo'laklarga bo'lish."""
    bolak, joriy = [], ""
    for q in qatorlar:
        if len(joriy) + len(q) + 1 > limit:
            bolak.append(joriy)
            joriy = ""
        joriy += q + "\n"
    if joriy.strip():
        bolak.append(joriy)
    return bolak


# --- Excel ------------------------------------------------------------------------
_SARL = Font(name="Arial", bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F4E78")
_ODDIY = Font(name="Arial")
_QALIN = Font(name="Arial", bold=True)
_CHIZIQ = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
_SON = "#,##0"


def _varaq(wb, nom, sarlavha, qatorlar, kengliklar, jami=None):
    ws = wb.create_sheet(nom)
    ws.append(sarlavha)
    for c in ws[1]:
        c.font, c.fill = _SARL, _FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for q in qatorlar:
        ws.append(q)
    if jami:
        ws.append(jami)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            # Jami qatori qalin, qolganlari oddiy
            c.font = _QALIN if (jami and c.row == ws.max_row) else _ODDIY
            c.border = _CHIZIQ
            if isinstance(c.value, (int, float)):
                c.number_format = _SON
    for i, k in enumerate(kengliklar):
        ws.column_dimensions[chr(65 + i)].width = k
    ws.freeze_panes = "A2"
    return ws


def kunlik_excel(sana, yol):
    """Kunlik hisobotni Excel faylga yozadi (qiymatlar — kun holatining surati)."""
    s = sana.isoformat()
    wb = Workbook()
    wb.remove(wb.active)

    f = db.oxirgi_fayl(s, "sotuv")
    q = db.fayl_qatorlari(f) if f else []
    fi = db.oxirgi_fayl(s, "ichki")
    # 0 so'mlik va ichki chiqimlar alohida varaqda
    narxsiz = [x for x in q if x["sotilgan"] > 0 and x["daromad"] == 0] + \
        (db.fayl_qatorlari(fi) if fi else [])
    q = [x for x in q if not (x["sotilgan"] > 0 and x["daromad"] == 0)]
    gr = ainur.modellarga_guruhla(q)
    _varaq(wb, "Modellar", ["Model", "Dona", "O'rtacha narx", "Summa", "Sotuvlar soni",
                            "0 so'mga dona"],
           [[g["model"], g["dona"], round(g["summa"] / g["dona"]) if g["dona"] else 0,
             g["summa"], g["sotish"], g["narxsiz_dona"]] for g in gr],
           [28, 10, 14, 16, 12, 12],
           ["JAMI", sum(g["dona"] for g in gr), None, sum(g["summa"] for g in gr),
            sum(g["sotish"] for g in gr), sum(g["narxsiz_dona"] for g in gr)])
    _varaq(wb, "Pozitsiyalar", ["Nomi", "Shtrix-kod", "Sotilgan", "Daromad", "Tannarx",
                                "Foyda", "O'rtacha narx"],
           [[x["nomi"], x["shtrix"], x["sotilgan"], x["daromad"], x["tannarx"], x["foyda"],
             round(x["daromad"] / x["sotilgan"]) if x["sotilgan"] else 0]
            for x in sorted(q, key=lambda x: -x["sotilgan"])],
           [30, 16, 10, 16, 16, 16, 14])

    if narxsiz:
        _varaq(wb, "NarxsizChiqim", ["Nomi", "Dona", "Summa"],
               [[x["nomi"], x["sotilgan"], x["daromad"]] for x in sorted(narxsiz, key=lambda x: -x["sotilgan"])],
               [30, 10, 14], ["JAMI", sum(x["sotilgan"] for x in narxsiz), sum(x["daromad"] for x in narxsiz)])

    hj = db.fayl_hujjatlari(f) if f else []
    if hj:
        turlar = list(dict.fromkeys(t for h in hj for t in h["turlar"]))
        _varaq(wb, "Hujjatlar", ["Raqam", "Vaqt", "Mijoz", "Sotuvchi", "Jami", "To'landi",
                                 "Qarzga"] + turlar,
               [[h["raqam"], h["vaqt"], h["mijoz"], h["xodim"], h["jami"], h["tolangan"],
                 max(h["jami"] - h["tolangan"], 0)] + [h["turlar"].get(t, 0) for t in turlar]
                for h in hj],
               [10, 8, 26, 16, 16, 16, 16] + [14] * len(turlar),
               ["JAMI", None, None, None, sum(h["jami"] for h in hj),
                sum(h["tolangan"] for h in hj),
                sum(max(h["jami"] - h["tolangan"], 0) for h in hj)]
               + [sum(h["turlar"].get(t, 0) for h in hj) for t in turlar])

    qt = db.kun_qarz_tolovlari(s)
    if qt:
        turlar = list(dict.fromkeys(t for x in qt for t in x["turlar"]))
        _varaq(wb, "QarzTolovlari", ["Mijoz", "Summa", "Hujjat", "Hujjat sanasi"] + turlar,
               [[x["mijoz"], x["summa"], x["raqam"], x["hujjat_sana"]]
                + [x["turlar"].get(t, 0) for t in turlar] for x in qt],
               [28, 16, 10, 14] + [14] * len(turlar),
               ["JAMI", sum(x["summa"] for x in qt), None, None]
               + [sum(x["turlar"].get(t, 0) for x in qt) for t in turlar])

    rx = db.rasxodlar(s)
    if rx:
        _varaq(wb, "Rasxodlar", ["Vaqt", "Kategoriya", "Summa", "Izoh", "Kim"],
               [[x["vaqt"], x["kategoriya"], x["summa"], x["izoh"], x["kim"]] for x in rx],
               [10, 20, 16, 40, 18], ["JAMI", None, sum(x["summa"] for x in rx), None, None])

    qz = sorted([m for m in db.qarz_ol(s) if m["qarz"] > 0], key=lambda m: -m["qarz"])
    if qz:
        _varaq(wb, "Qarzdorlik", ["Mijoz", "Qarz"], [[m["nom"], m["qarz"]] for m in qz],
               [30, 18], ["JAMI", sum(m["qarz"] for m in qz)])

    f = db.oxirgi_fayl(s, "qoldiq")
    if f:
        oq = db.fayl_qatorlari(f)
        _varaq(wb, "Ombor", ["Nomi", "Shtrix-kod", "Boshlang'ich", "Kelish", "Chiqim",
                             "Qoldiq"],
               [[x["nomi"], x["shtrix"], x["boshl"], x["kelish"], x["istemol"], x["yakuniy"]]
                for x in sorted(oq, key=lambda x: x["nomi"])],
               [30, 16, 13, 11, 11, 11],
               ["JAMI", None, sum(x["boshl"] for x in oq), sum(x["kelish"] for x in oq),
                sum(x["istemol"] for x in oq), sum(x["yakuniy"] for x in oq)])
    wb.save(yol)
    return yol
